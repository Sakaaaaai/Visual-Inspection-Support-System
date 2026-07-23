from __future__ import annotations

import os
import re
import shutil
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ANIMAL_OPTIONS = (
    "bear（クマ）",
    "boar（イノシシ）",
    "racoondog（タヌキ）",
    "man（ヒト）",
    "car（クルマ）",
    "maskedmusang（ハクビシン）",
    "dog（イヌ）",
    "cat（ネコ）",
    "deer（シカ）",
    "fox（キツネ）",
    "serow（カモシカ）",
    "rabbit（ウサギ）",
    "craw（カラス）",
    "monkey（サル）",
    "badger（アナグマ）",
    "racoon（アライグマ）",
    "?",
    "いない",
)

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".gif", ".webp"}


class WorkbookFormatError(ValueError):
    pass


@dataclass(frozen=True)
class ColumnMap:
    filename: int
    predicted_animal: int
    predicted_count: int
    device: int
    manual_animal: int
    manual_count: int
    ai_animal: int
    ai_confidence: int


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _normalized(value: Any) -> str:
    return re.sub(r"[\s　:_：]+", "", _text(value)).lower()


def animal_code(value: Any) -> str:
    """Return the fixed folder code at the start of an animal label."""
    match = re.match(r"\s*([A-Za-z]+)", _text(value))
    if not match:
        raise WorkbookFormatError(f"動物名から英語コードを取得できません: {_text(value) or '（空欄）'}")
    return match.group(1).lower()


def device_folder(value: Any) -> str:
    normalized = _normalized(value)
    if "server" in normalized or "サーバ" in normalized:
        return "server"
    if "device" in normalized or "装置" in normalized:
        return "device"
    raise WorkbookFormatError(f"デバイスを server/device に変換できません: {_text(value) or '（空欄）'}")


def get_status(manual_animal: Any, manual_count: Any) -> str:
    text_animal = _text(manual_animal)
    if text_animal == "保留":
        return "hold"
    if text_animal == "いない" or manual_count not in (None, ""):
        return "reviewed"
    return "pending"


def sequence_key_and_frame(filename: Any) -> tuple[str, int] | None:
    """Split a capture filename like ``...1939_4.jpeg`` into its event and frame."""
    match = re.match(r"^(.*)_(\d+)$", Path(_text(filename)).stem)
    if not match:
        return None
    return match.group(1), int(match.group(2))


def capture_sort_key(filename: Any) -> tuple[datetime, int] | None:
    """Read capture time and frame from ``..._2026.06.12.1939_4.jpeg``."""
    stem = Path(_text(filename)).stem
    match = re.search(r"_(\d{4}\.\d{2}\.\d{2}\.\d{4})_(\d+)$", stem)
    if not match:
        return None
    try:
        captured_at = datetime.strptime(match.group(1), "%Y.%m.%d.%H%M")
    except ValueError:
        return None
    return captured_at, int(match.group(2))


def normalized_manual_values(predicted_animal: str, selected_animal: str, count: Any) -> tuple[str | None, int | None]:
    if selected_animal == "保留":
        return "保留", None
    if selected_animal not in ANIMAL_OPTIONS:
        raise ValueError("動物名は候補から選択してください。")
    if selected_animal == "いない":
        return "いない", None

    if isinstance(count, bool):
        raise ValueError("数は0以上の整数で入力してください。")
    try:
        parsed_count = int(str(count).strip())
    except (TypeError, ValueError):
        raise ValueError("数は0以上の整数で入力してください。") from None
    if str(parsed_count) != str(count).strip() or parsed_count < 0:
        raise ValueError("数は0以上の整数で入力してください。")

    manual_animal = None if selected_animal == predicted_animal else selected_animal
    return manual_animal, parsed_count


class AnimalWorkbook:
    def __init__(self, batch_dir: Path, workbook_path: Path):
        self.batch_dir = batch_dir.resolve()
        self.path = workbook_path.resolve()
        self.keep_vba = self.path.suffix.lower() == ".xlsm"
        self.workbook = load_workbook(self.path, keep_vba=self.keep_vba, data_only=False)
        self.sheet, self.header_row, self.columns = self._find_data_sheet()
        self.device_number, self.analysis_date = self._read_metadata()
        self.rows = self._find_image_rows()
        self._image_catalog = self._build_image_catalog()
        self.backup_path: Path | None = None

        if not self.rows:
            raise WorkbookFormatError("対象Excelに画像ファイルの行が見つかりません。")

    @classmethod
    def open_batch(cls, batch_dir: str | os.PathLike[str]) -> "AnimalWorkbook":
        root = Path(batch_dir).expanduser().resolve()
        if not root.is_dir():
            raise FileNotFoundError(f"解析結果フォルダが見つかりません: {root}")
        candidates = sorted(
            path for path in root.glob("result_*_all.xlsx")
            if not path.name.startswith("~$")
        )
        candidates += sorted(
            path for path in root.glob("result_*_all.xlsm")
            if not path.name.startswith("~$")
        )
        if len(candidates) != 1:
            raise WorkbookFormatError(
                f"result_*_all.xlsx が1つ必要です。見つかった数: {len(candidates)}"
            )
        return cls(root, candidates[0])

    def close(self) -> None:
        self.workbook.close()

    def _find_data_sheet(self):
        required = {
            "filename": ("ファイル名",),
            "predicted_animal": ("動物名",),
            "predicted_count": ("数",),
            "device": ("デバイス",),
            "manual_animal": ("目視による動物名", "目視動物名"),
            "manual_count": ("目視による数", "目視数"),
        }
        for sheet in self.workbook.worksheets:
            for row_number in range(1, min(sheet.max_row, 30) + 1):
                headers = {
                    _normalized(sheet.cell(row_number, column).value): column
                    for column in range(1, sheet.max_column + 1)
                    if _text(sheet.cell(row_number, column).value)
                }
                found: dict[str, int] = {}
                for key, aliases in required.items():
                    for alias in aliases:
                        column = headers.get(_normalized(alias))
                        if column:
                            found[key] = column
                            break
                if len(found) == len(required):
                    ai_animal, ai_confidence = self._ensure_ai_columns(sheet, row_number, headers)
                    found["ai_animal"] = ai_animal
                    found["ai_confidence"] = ai_confidence
                    return sheet, row_number, ColumnMap(**found)
        raise WorkbookFormatError(
            "必要な列（ファイル名、動物名、数、デバイス、目視による動物名、目視による数）が見つかりません。"
        )

    @staticmethod
    def _ensure_ai_columns(sheet, row_number: int, headers: dict[str, int]) -> tuple[int, int]:
        """Find or create the columns used to store CLIP-based AI re-predictions."""
        ai_animal = headers.get(_normalized("AI再判定"))
        ai_confidence = headers.get(_normalized("AI確信度"))
        next_column = sheet.max_column + 1
        if ai_animal is None:
            ai_animal = next_column
            sheet.cell(row_number, ai_animal).value = "AI再判定"
            next_column += 1
        if ai_confidence is None:
            ai_confidence = next_column
            sheet.cell(row_number, ai_confidence).value = "AI確信度"
        return ai_animal, ai_confidence

    def _read_metadata(self) -> tuple[str, str]:
        values = []
        for row in self.sheet.iter_rows(min_row=1, max_row=min(self.header_row, 10), values_only=True):
            values.extend(_text(value) for value in row if value is not None)
        joined = " | ".join(values)
        device_match = re.search(r"装置番号\s*[:：]\s*(\d+)", joined)
        date_match = re.search(r"解析日\s*[:：]\s*(\d{6,8})", joined)

        filename_match = re.match(r"result_(\d+)_(\d{6,8})_all\.(?:xlsx|xlsm)$", self.path.name, re.I)
        device = device_match.group(1) if device_match else (filename_match.group(1) if filename_match else "")
        analysis_date = date_match.group(1) if date_match else (filename_match.group(2) if filename_match else "")
        if not device or not analysis_date:
            raise WorkbookFormatError("装置番号または解析日を取得できません。")
        return device, analysis_date

    def _find_image_rows(self) -> list[int]:
        result = []
        for row in range(self.header_row + 1, self.sheet.max_row + 1):
            filename = _text(self.sheet.cell(row, self.columns.filename).value)
            if Path(filename).suffix.lower() in IMAGE_EXTENSIONS:
                result.append(row)
        return result

    def _build_image_catalog(self) -> dict[str, list[tuple[tuple[datetime, int], str, Path, bool]]]:
        """Index every image, including undetected files stored at the device root."""
        catalogs: dict[str, list[tuple[tuple[datetime, int], str, Path, bool]]] = {}
        extension_priority = {".jpg": 0, ".jpeg": 1, ".png": 2, ".bmp": 3, ".webp": 4, ".gif": 5}
        for device_name in ("server", "device"):
            device_root = self.batch_dir / device_name
            if not device_root.is_dir():
                catalogs[device_name] = []
                continue
            best_by_identity: dict[str, tuple[tuple[int, int, str], Path, tuple[datetime, int], bool]] = {}
            for path in device_root.rglob("*"):
                if not path.is_file() or path.suffix.lower() not in IMAGE_EXTENSIONS:
                    continue
                sort_key = capture_sort_key(path.name)
                if sort_key is None:
                    continue
                identity = path.stem
                parent_name = path.parent.name.lower()
                if "_inf_" in parent_name:
                    location_priority = 0
                    detected = True
                elif "_org_" in parent_name:
                    location_priority = 1
                    detected = True
                elif path.parent == device_root:
                    location_priority = 2
                    detected = False
                else:
                    location_priority = 3
                    detected = False
                priority = (
                    location_priority,
                    extension_priority.get(path.suffix.lower(), 9),
                    str(path).lower(),
                )
                existing = best_by_identity.get(identity)
                if existing is None or priority < existing[0]:
                    best_by_identity[identity] = (priority, path, sort_key, detected)
            catalogs[device_name] = sorted(
                (
                    (sort_key, identity, path, detected)
                    for identity, (_, path, sort_key, detected) in best_by_identity.items()
                ),
                key=lambda item: (*item[0], item[1]),
            )
        return catalogs

    def image_candidates_for_row(self, excel_row: int) -> tuple[Path, ...]:
        filename = _text(self.sheet.cell(excel_row, self.columns.filename).value)
        animal = animal_code(self.sheet.cell(excel_row, self.columns.predicted_animal).value)
        device = device_folder(self.sheet.cell(excel_row, self.columns.device).value)
        base = self.batch_dir / device
        inf_folder = f"{self.device_number}_{self.analysis_date}_inf_{animal}"
        org_folder = f"{self.device_number}_{self.analysis_date}_org_{animal}"
        source = Path(filename)
        alternate_suffixes = [suffix for suffix in (".jpg", ".jpeg", ".png") if suffix != source.suffix.lower()]
        names = [source.name, *(f"{source.stem}{suffix}" for suffix in alternate_suffixes)]
        return tuple(
            base / folder / name
            for folder in (inf_folder, org_folder)
            for name in names
        )

    def image_path_for_row(self, excel_row: int) -> Path:
        candidates = self.image_candidates_for_row(excel_row)
        for candidate in candidates:
            if candidate.is_file():
                return candidate
        return candidates[0]

    def _context_entries(
        self, index: int, radius: int = 2
    ) -> list[tuple[tuple[datetime, int], str, Path, bool]]:
        excel_row = self.rows[index]
        current_filename = _text(self.sheet.cell(excel_row, self.columns.filename).value)
        current_sort_key = capture_sort_key(current_filename)
        if current_sort_key is None:
            return []
        current_identity = Path(current_filename).stem
        try:
            current_device = device_folder(
                self.sheet.cell(excel_row, self.columns.device).value
            )
        except WorkbookFormatError:
            return []
        catalog = self._image_catalog.get(current_device, [])
        current_position = next(
            (position for position, item in enumerate(catalog) if item[1] == current_identity),
            None,
        )
        if current_position is None:
            return []
        return catalog[max(0, current_position - radius): current_position + radius + 1]

    def context_images(self, index: int, radius: int = 2) -> list[dict[str, Any]]:
        """Return nearby detected and undetected images without changing the review row."""
        current_filename = _text(
            self.sheet.cell(self.rows[index], self.columns.filename).value
        )
        current_identity = Path(current_filename).stem
        entries = self._context_entries(index, radius)
        return [
            {
                "contextPosition": position,
                "frame": sort_key[1],
                "capturedAt": sort_key[0].strftime("%Y-%m-%d %H:%M"),
                "filename": path.name,
                "imagePath": str(path),
                "imageExists": path.is_file(),
                "isCurrent": identity == current_identity,
                "detected": detected,
            }
            for position, (sort_key, identity, path, detected) in enumerate(entries)
        ]

    def context_image_path(self, index: int, context_position: int) -> Path:
        entries = self._context_entries(index)
        if context_position < 0 or context_position >= len(entries):
            raise IndexError("参考画像が範囲外です。")
        return entries[context_position][2]

    def row_data(self, index: int) -> dict[str, Any]:
        excel_row = self.rows[index]
        predicted_animal = _text(self.sheet.cell(excel_row, self.columns.predicted_animal).value)
        predicted_count = self.sheet.cell(excel_row, self.columns.predicted_count).value
        manual_animal = _text(self.sheet.cell(excel_row, self.columns.manual_animal).value)
        manual_count = self.sheet.cell(excel_row, self.columns.manual_count).value
        path_error = None
        try:
            image_path = self.image_path_for_row(excel_row)
            if not image_path.is_file():
                candidates = self.image_candidates_for_row(excel_row)
                path_error = "画像が見つかりません。確認先: " + " / ".join(map(str, candidates))
        except WorkbookFormatError as exc:
            image_path = None
            path_error = str(exc)

        status = get_status(manual_animal, manual_count)
        ui_selected_animal = predicted_animal if status == "hold" else (manual_animal or predicted_animal)
        displayed_count = "" if ui_selected_animal in ("いない", "保留") else (
            manual_count if manual_count not in (None, "") else predicted_count
        )
        ai_animal = _text(self.sheet.cell(excel_row, self.columns.ai_animal).value)
        ai_confidence = self.sheet.cell(excel_row, self.columns.ai_confidence).value
        return {
            "index": index,
            "excelRow": excel_row,
            "filename": _text(self.sheet.cell(excel_row, self.columns.filename).value),
            "predictedAnimal": predicted_animal,
            "predictedCount": predicted_count,
            "device": _text(self.sheet.cell(excel_row, self.columns.device).value),
            "selectedAnimal": ui_selected_animal,
            "manualCount": displayed_count,
            "status": status,
            "imageExists": bool(image_path and image_path.is_file()),
            "imagePath": str(image_path) if image_path else "",
            "pathError": path_error,
            "contextImages": self.context_images(index),
            "aiAnimal": ai_animal,
            "aiConfidence": ai_confidence,
        }

    def completed_count(self) -> int:
        return sum(
            1 for row in self.rows
            if get_status(
                self.sheet.cell(row, self.columns.manual_animal).value,
                self.sheet.cell(row, self.columns.manual_count).value,
            ) == "reviewed"
        )

    def next_with_status(self, current_index: int, target_statuses: list[str]) -> int | None:
        total = len(self.rows)
        for offset in range(1, total + 1):
            index = (current_index + offset) % total
            row = self.rows[index]
            status = get_status(
                self.sheet.cell(row, self.columns.manual_animal).value,
                self.sheet.cell(row, self.columns.manual_count).value,
            )
            if status in target_statuses:
                return index
        return None

    def rows_by_animal(self) -> dict[str, list[int]]:
        """Group row indices by predicted animal name."""
        groups: dict[str, list[int]] = {}
        for index, excel_row in enumerate(self.rows):
            animal = _text(self.sheet.cell(excel_row, self.columns.predicted_animal).value)
            if animal not in groups:
                groups[animal] = []
            groups[animal].append(index)
        return groups

    def grid_row_data(self, index: int) -> dict[str, Any]:
        """Return minimal row data for grid display (no context images)."""
        excel_row = self.rows[index]
        predicted_animal = _text(self.sheet.cell(excel_row, self.columns.predicted_animal).value)
        predicted_count = self.sheet.cell(excel_row, self.columns.predicted_count).value
        manual_animal = _text(self.sheet.cell(excel_row, self.columns.manual_animal).value)
        manual_count = self.sheet.cell(excel_row, self.columns.manual_count).value
        path_error = None
        try:
            image_path = self.image_path_for_row(excel_row)
            if not image_path.is_file():
                path_error = "画像が見つかりません"
        except WorkbookFormatError as exc:
            image_path = None
            path_error = str(exc)
        status = get_status(manual_animal, manual_count)
        return {
            "index": index,
            "filename": _text(self.sheet.cell(excel_row, self.columns.filename).value),
            "predictedAnimal": predicted_animal,
            "predictedCount": predicted_count,
            "status": status,
            "imageExists": bool(image_path and image_path.is_file()),
            "pathError": path_error,
        }

    def confirm_prediction(self, index: int) -> None:
        """Save the AI prediction as the manual review result (prediction is correct)."""
        if index < 0 or index >= len(self.rows):
            raise IndexError("対象行が範囲外です。")
        row = self.rows[index]
        predicted_animal = _text(self.sheet.cell(row, self.columns.predicted_animal).value)
        predicted_count = self.sheet.cell(row, self.columns.predicted_count).value
        if predicted_animal == "いない":
            self.sheet.cell(row, self.columns.manual_animal).value = "いない"
            self.sheet.cell(row, self.columns.manual_count).value = None
        else:
            # When manual matches predicted, manual_animal stays None
            self.sheet.cell(row, self.columns.manual_animal).value = None
            try:
                self.sheet.cell(row, self.columns.manual_count).value = int(predicted_count)
            except (TypeError, ValueError):
                self.sheet.cell(row, self.columns.manual_count).value = predicted_count

    def confirm_predictions_bulk(self, indices: list[int]) -> int:
        """Confirm multiple predictions at once, saving atomically. Returns count confirmed."""
        previous_values: list[tuple[int, Any, Any]] = []
        for index in indices:
            if index < 0 or index >= len(self.rows):
                continue
            row = self.rows[index]
            animal_cell = self.sheet.cell(row, self.columns.manual_animal)
            count_cell = self.sheet.cell(row, self.columns.manual_count)
            previous_values.append((row, animal_cell.value, count_cell.value))
            self.confirm_prediction(index)
        try:
            self._save_atomic()
        except Exception:
            for row, old_animal, old_count in previous_values:
                self.sheet.cell(row, self.columns.manual_animal).value = old_animal
                self.sheet.cell(row, self.columns.manual_count).value = old_count
            raise
        return len(previous_values)

    def save_review_bulk(self, indices: list[int], selected_animal: str, count: Any) -> int:
        """Override multiple predictions at once with the specified animal and count."""
        previous_values: list[tuple[int, Any, Any]] = []
        for index in indices:
            if index < 0 or index >= len(self.rows):
                continue
            row = self.rows[index]
            predicted = _text(self.sheet.cell(row, self.columns.predicted_animal).value)
            manual_animal, manual_count = normalized_manual_values(predicted, selected_animal, count)
            
            animal_cell = self.sheet.cell(row, self.columns.manual_animal)
            count_cell = self.sheet.cell(row, self.columns.manual_count)
            previous_values.append((row, animal_cell.value, count_cell.value))
            
            animal_cell.value = manual_animal
            count_cell.value = manual_count

        try:
            self._save_atomic()
        except Exception:
            for row, old_animal, old_count in previous_values:
                self.sheet.cell(row, self.columns.manual_animal).value = old_animal
                self.sheet.cell(row, self.columns.manual_count).value = old_count
            raise
        return len(previous_values)

    def save_review(self, index: int, selected_animal: str, count: Any) -> dict[str, Any]:
        if index < 0 or index >= len(self.rows):
            raise IndexError("対象行が範囲外です。")
        row = self.rows[index]
        predicted = _text(self.sheet.cell(row, self.columns.predicted_animal).value)
        manual_animal, manual_count = normalized_manual_values(predicted, selected_animal, count)

        animal_cell = self.sheet.cell(row, self.columns.manual_animal)
        count_cell = self.sheet.cell(row, self.columns.manual_count)
        previous = (animal_cell.value, count_cell.value)
        animal_cell.value = manual_animal
        count_cell.value = manual_count
        try:
            self._save_atomic()
        except Exception:
            animal_cell.value, count_cell.value = previous
            raise
        return self.row_data(index)

    def save_ai_prediction(self, index: int, animal: str, confidence: float) -> dict[str, Any]:
        """Store a CLIP-based AI re-prediction for a row, separate from the device/server prediction."""
        if index < 0 or index >= len(self.rows):
            raise IndexError("対象行が範囲外です。")
        row = self.rows[index]
        animal_cell = self.sheet.cell(row, self.columns.ai_animal)
        confidence_cell = self.sheet.cell(row, self.columns.ai_confidence)
        previous = (animal_cell.value, confidence_cell.value)
        animal_cell.value = animal
        confidence_cell.value = round(confidence, 4)
        try:
            self._save_atomic()
        except Exception:
            animal_cell.value, confidence_cell.value = previous
            raise
        return self.row_data(index)

    def _save_atomic(self) -> None:
        if self.backup_path is None:
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            backup = self.path.with_name(f"{self.path.stem}_backup_{stamp}{self.path.suffix}")
            shutil.copy2(self.path, backup)
            self.backup_path = backup

        fd, temporary_name = tempfile.mkstemp(
            prefix=f".{self.path.stem}_",
            suffix=self.path.suffix,
            dir=self.path.parent,
        )
        os.close(fd)
        temporary = Path(temporary_name)
        try:
            self.workbook.save(temporary)
            os.replace(temporary, self.path)
        finally:
            if temporary.exists():
                temporary.unlink()
